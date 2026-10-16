"""Module 1 and 2: resume upload validation and text extraction (PDF and DOCX)."""
import os

from pypdf import PdfReader
import docx  # python-docx

MAX_MB = 5
ALLOWED = {".pdf", ".docx"}


def validate_file(filename: str, size_bytes: int):
    """Return (ok, message). Checks the file type and the file size."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED:
        return False, "Only PDF and DOCX files are supported."
    if size_bytes > MAX_MB * 1024 * 1024:
        return False, f"File is too large. Maximum size is {MAX_MB} MB."
    return True, "OK"


def extract_pdf(file_like) -> str:
    """Extract text from every PDF page."""
    reader = PdfReader(file_like)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_docx(file_like) -> str:
    """Extract text from every DOCX paragraph and table cell."""
    document = docx.Document(file_like)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(file_like, filename: str) -> str:
    """Pick the right extractor from the file extension."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_pdf(file_like)
    if ext == ".docx":
        return extract_docx(file_like)
    raise ValueError("Unsupported file type: " + ext)
